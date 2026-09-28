import os
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from decimal import Decimal

from app.db.database import SessionLocal, engine, Base
# Import all models so Base.metadata.create_all works
from app.models import pan_master, supplier, purchase_order, invoice, gcp_invoice

Base.metadata.create_all(bind=engine)

EXCEL_FILE = Path("data/gcp_invoice_data.xlsx")
SHEET_NAME = "in"

def clean_date(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None

def clean_decimal(val):
    if val is None or val == "":
        return None
    try:
        return Decimal(str(val))
    except:
        return None

def clean_int(val):
    if val is None or val == "":
        return None
    try:
        return int(val)
    except:
        return None

def clean_str(val):
    if val is None:
        return None
    val = str(val).strip()
    return val if val != "" else None

def main():
    print("Loading workbook...")
    workbook = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    
    db = SessionLocal()
    
    print("Clearing old gcp_invoice records...")
    db.query(gcp_invoice.GCPInvoice).delete()
    db.commit()
    
    print("Inserting records...")
    records = []
    for row in rows:
        record = dict(zip(headers, row))
        obj = gcp_invoice.GCPInvoice(
            inv_no=clean_str(record.get("INV_NO")),
            inv_dt=clean_date(record.get("INV_DT")),
            doc_type=clean_str(record.get("DOC_TYPE")),
            status_flag=clean_str(record.get("STATUS_FLAG")),
            overall_status=clean_str(record.get("OVERALL_STATUS")),
            po_no=clean_str(record.get("PO_NO")),
            po_item=clean_int(record.get("PO_ITEM")),
            supplier=clean_str(record.get("SUPPLIER")),
            supplier_name=clean_str(record.get("SUPPLIER_NAME")),
            pan_no=clean_str(record.get("PAN_NO")),
            approved_by=clean_str(record.get("APPROVED_BY")),
            pending_with=clean_str(record.get("PENDING_WITH")),
            approver1=clean_str(record.get("APPROVER1")),
            approver2=clean_str(record.get("APPROVER2")),
            approver3=clean_str(record.get("APPROVER3")),
            approver4=clean_str(record.get("APPROVER4")),
            approver5=clean_str(record.get("APPROVER5")),
            approver6=clean_str(record.get("APPROVER6")),
            approver7=clean_str(record.get("APPROVER7")),
            wf_levels=clean_int(record.get("WF_LEVELS")),
            c_wf_level=clean_int(record.get("C_WF_LEVEL")),
            wf_steps=clean_str(record.get("WF_STEPS")),
            rej_reason=clean_str(record.get("REJ_REASON")),
            accountant=clean_str(record.get("ACCOUNTANT")),
            doc_sent_acc=clean_date(record.get("DOC_SENT_ACC")),
            final_appr_dt=clean_date(record.get("FINAL_APPR_DT")),
            appr_remarks=clean_str(record.get("APPR_REMARKS")),
            miro_document=clean_str(record.get("MIRO_DOCUMENT")),
            miro_company_code=clean_str(record.get("MIRO_COMPANY_CODE")),
            miro_fiscal_year=clean_int(record.get("MIRO_FISCAL_YEAR")),
            miro_document_type=clean_str(record.get("MIRO_DOCUMENT_TYPE")),
            miro_document_date=clean_date(record.get("MIRO_DOCUMENT_DATE")),
            miro_posting_date=clean_date(record.get("MIRO_POSTING_DATE")),
            miro_reference=clean_str(record.get("MIRO_REFERENCE")),
            miro_currency=clean_str(record.get("MIRO_CURRENCY")),
            miro_amount=clean_decimal(record.get("MIRO_AMOUNT")),
            miro_amount_lc=clean_decimal(record.get("MIRO_AMOUNT_LC")),
            miro_cd_base=clean_decimal(record.get("MIRO_CD_BASE")),
            miro_po_no=clean_str(record.get("MIRO_PO_NO")),
            miro_po_item=clean_int(record.get("MIRO_PO_ITEM")),
            miro_net_due_date=clean_date(record.get("MIRO_NET_DUE_DATE")),
            miro_baseline_date=clean_date(record.get("MIRO_BASELINE_DATE")),
            miro_account_type=clean_str(record.get("MIRO_ACCOUNT_TYPE")),
            miro_posting_key=clean_str(record.get("MIRO_POSTING_KEY")),
            miro_clearing_date=clean_date(record.get("MIRO_CLEARING_DATE")),
            miro_clearing_document=clean_str(record.get("MIRO_CLEARING_DOCUMENT")),
            miro_clearing_fiscal_year=clean_int(record.get("MIRO_CLEARING_FISCAL_YEAR")),
            miro_clearing_status=clean_str(record.get("MIRO_CLEARING_STATUS")),
            utr_no=clean_str(record.get("UTR_NO"))
        )
        records.append(obj)
        
        if len(records) >= 500:
            db.bulk_save_objects(records)
            db.commit()
            records = []
            
    if records:
        db.bulk_save_objects(records)
        db.commit()
        
    print("Done!")
    workbook.close()
    db.close()

if __name__ == "__main__":
    main()
