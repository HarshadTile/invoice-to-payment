import os

import pytest

from otp_tests.sample_data import make_mobile
from fastapi import HTTPException
from openpyxl import Workbook

from app.repositories import gcp_invoice
from app.services.supplier_mobile import registered_mobile


def write_book(path, phones):
    book = Workbook()
    sheet = book.active
    sheet.title = "in"
    sheet.append(["SUPPLIER", "SUPPLIER_MOBILE", "MIRO_CD_BASE"])
    for phone in phones:
        sheet.append(["TEST-VENDOR", phone, 100])
        sheet.cell(sheet.max_row, 2).number_format = "@"
    book.save(path)
    book.close()
    # Deterministic file-change signal, including files saved rapidly in tests.
    stat = path.stat()
    os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))


@pytest.fixture
def workbook(tmp_path, monkeypatch):
    monkeypatch.delenv("SUPPLIER_MOBILE_EXCEL_FILE", raising=False)
    path = tmp_path / "test-invoices.xlsx"
    monkeypatch.setattr(gcp_invoice, "EXCEL_FILE", path)
    monkeypatch.setattr(gcp_invoice, "_cached_invoices", None)
    monkeypatch.setattr(gcp_invoice, "_cached_excel_signature", None)
    return path


def test_saved_mobile_change_refreshes_cache(workbook):
    write_book(workbook, [make_mobile("old-phone"), make_mobile("old-phone")])
    first = gcp_invoice.get_all_cached_invoices()
    assert registered_mobile("TEST-VENDOR", first) == make_mobile("old-phone")
    assert gcp_invoice.get_all_cached_invoices() is first
    write_book(workbook, [make_mobile("new-phone"), make_mobile("new-phone")])
    second = gcp_invoice.get_all_cached_invoices()
    assert second is not first
    assert registered_mobile("TEST-VENDOR", second) == make_mobile("new-phone")
    assert second[0]["MIRO_AMOUNT"] == 100


def test_partial_vendor_update_is_rejected_after_reload(workbook):
    write_book(workbook, [make_mobile("old-phone"), make_mobile("old-phone")])
    gcp_invoice.get_all_cached_invoices()
    write_book(workbook, [make_mobile("old-phone"), make_mobile("new-phone")])
    with pytest.raises(HTTPException) as error:
        registered_mobile("TEST-VENDOR", gcp_invoice.get_all_cached_invoices())
    assert error.value.status_code == 400


def test_changed_workbook_path_refreshes_even_with_same_stat(workbook, monkeypatch):
    write_book(workbook, [make_mobile("old-phone")])
    first = gcp_invoice.get_all_cached_invoices()
    second_path = workbook.parent / "another.xlsx"
    write_book(second_path, [make_mobile("new-phone")])
    monkeypatch.setattr(gcp_invoice, "EXCEL_FILE", second_path)
    assert gcp_invoice.get_all_cached_invoices() is not first
    assert registered_mobile("TEST-VENDOR", gcp_invoice.get_all_cached_invoices()) == make_mobile("new-phone")


def test_missing_workbook_does_not_fall_back_to_cached_mobile(workbook):
    write_book(workbook, [make_mobile("old-phone")])
    gcp_invoice.get_all_cached_invoices()
    workbook.unlink()
    with pytest.raises(FileNotFoundError):
        gcp_invoice.get_all_cached_invoices()


def test_workbook_changing_during_read_is_not_published(workbook, monkeypatch):
    write_book(workbook, [make_mobile("old-phone")])
    values = iter(range(10))
    monkeypatch.setattr(gcp_invoice, "_excel_signature", lambda: next(values))
    with pytest.raises(OSError, match="changed while reading"):
        gcp_invoice.get_all_cached_invoices()
    assert gcp_invoice._cached_invoices is None
