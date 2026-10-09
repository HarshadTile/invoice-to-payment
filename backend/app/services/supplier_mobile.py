"""Replace this adapter when the vendor master becomes available."""
import re
import os
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook


def _mapping_rows(path: str) -> list[dict]:
    """Small vendor master: read saved changes on each login/resend."""
    workbook = None
    try:
        source = Path(path)
        before = source.stat()
        workbook = load_workbook(source, read_only=True, data_only=True)
        rows = workbook["vendors"].iter_rows(values_only=True)
        headers = next(rows)
        if any(headers.count(column) != 1 for column in ("SUPPLIER", "SUPPLIER_MOBILE")):
            raise ValueError("Missing or duplicate vendor master columns")
        records = [dict(zip(headers, row)) for row in rows]
        after = source.stat()
        if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
            raise ValueError("Vendor master changed while reading")
        return records
    except Exception:
        raise HTTPException(503, "Supplier mobile workbook is unavailable or invalid. Contact support.") from None
    finally:
        if workbook is not None:
            workbook.close()


def registered_mobile(vcode: str, rows: list[dict]) -> str:
    mapping_file = os.getenv("SUPPLIER_MOBILE_EXCEL_FILE", "").strip()
    if mapping_file:
        rows = _mapping_rows(mapping_file)
    matches = [row for row in rows if str(row.get("SUPPLIER", "")).strip().lower() == vcode.lower()]
    if mapping_file and len(matches) > 1:
        raise HTTPException(400, "Duplicate vendor rows in the supplier mobile workbook. Contact support.")
    phones = set()
    for row in matches:
        value = row.get("SUPPLIER_MOBILE")
        if not isinstance(value, str) or not re.fullmatch(r"\+[1-9][0-9]{7,14}", value.strip()):
            raise HTTPException(400, "Registered mobile number is missing or invalid. Contact support.")
        phones.add(value.strip())
    if len(phones) != 1:
        raise HTTPException(400, "Registered mobile mapping is unavailable or conflicting. Contact support.")
    return phones.pop()
