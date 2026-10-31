"""SRM invoice source; GCP can later supply the same source-row contract."""
import re
from datetime import date, datetime
from pathlib import Path
from threading import Lock

from openpyxl import load_workbook
from app.repositories.workbook_cache import read_workbook_snapshot

EXCEL_FILE = Path(__file__).resolve().parents[2] / "data" / "tbl_dw_inv_paymt_SRM.xlsx"
SHEET_NAME = "in"
_cached_invoices = None
_load_lock = Lock()

STATUS_LABELS = {
    "D": "Rejected", "Y": "Rejected", "X": "Entry Deleted",
    "N": "Push to PO Portal", "L": "Level Approved", "A": "Approved",
    "P": "Pending for Approval", "I": "Initial",
}

DATE_FIELDS = {
    "invoice": ("INVOICE_DATE",), "upload": ("INVOICE_UPLOAD_DATE",),
    "approval": ("FINAL_APPRPVAL_DATE",), "miro": ("MIRO_DOCUMENT_DATE",),
    "due": ("NETDT", "NET_DUE_DATE", "Payment_Due_Date"), "payment": ("CLEARING_DATE",),
    "posting": ("POSTING_DATE",), "accounts": ("DOC_SENT_ACCOUNT_DATE",),
    "utr_clearing": ("UTR_CLEARING_DATE_PAID",),
}


def text_value(value):
    if value is None:
        return None
    value = str(value).strip()
    return None if value in ("", "-") else value


_DAY_MONTH_YEAR = re.compile(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\b")
# The source's UTR_CLEARING_DATE_PAID is often the UTR with the date glued on the end,
# e.g. "HDFCH008095981872026-02-17" -> 2026-02-17.
_TRAILING_ISO_DATE = re.compile(r"(\d{4}-\d{2}-\d{2})(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?$")


def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = text_value(value)
    if not value:
        return None
    candidates = [value[:10]]
    if match := _TRAILING_ISO_DATE.search(value):
        candidates.append(match.group(1))
    for candidate in candidates:
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            pass
    if match := _DAY_MONTH_YEAR.match(value):
        day, month, year = (int(part) for part in match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            pass
    return None


def invoice_status(row):
    """Payment evidence wins, then MIRO booking, then the SRM approval flag."""
    if row.get('__channel') == 'poPortal':
        from app.repositories.po_portal_invoice import invoice_status as po_status
        return po_status(row)
    if text_value(row.get("UTR_NUMBER")) or date_value(row.get("CLEARING_DATE")):
        return "Paid"
    if (text_value(row.get("MIRO_BOOKED")) or "").upper() == "MIRO BOOKED":
        return "Miro Booked"
    flag = (text_value(row.get("STATUS_FLAG")) or "").upper()
    if flag == "I" and (text_value(row.get("Invoice_Uploaded")) or "").upper() == "YES":
        return "Invoice Uploaded"
    return STATUS_LABELS.get(flag, flag or "Unknown")


def _load_excel_data():
    global _cached_invoices
    if _cached_invoices is not None:
        return _cached_invoices
    with _load_lock:
        if _cached_invoices is None:
            records = read_workbook_snapshot(EXCEL_FILE, SHEET_NAME, _read_excel_rows)
            for row in records:
                row["OVERALL_STATUS"] = invoice_status(row)
            _cached_invoices = records
    return _cached_invoices


def _read_excel_rows():
    workbook = load_workbook(EXCEL_FILE, read_only=True, data_only=True, keep_links=False)
    try:
        rows = workbook[SHEET_NAME].iter_rows(values_only=True)
        headers = next(rows)
        return [dict(zip(headers, values)) for values in rows if any(value is not None for value in values)]
    finally:
        workbook.close()


def filter_invoices(
    *, status=None, vendor_code=None, invoice_number=None, po_number=None,
    po_item=None, vendor_codes=None, product_type=None, plant=None, dv_type=None,
    upload_date_from=None, upload_date_to=None, filter_fiscal_year=None,
    date_type="invoice", date_from=None, date_to=None,
    approver=None, pending_with=None, approved_by=None, miro_status=None, payment_status=None,
    channel='msetuSrm', source_id=None, hod1_status=None, hod2_status=None,
    accounts_status=None, approval_stage=None, invoice_uploaded=None, po_type=None, exclude_deleted=False,
):
    allowed = {c.strip().lower() for c in vendor_codes} if vendor_codes is not None else None
    records = []
    sources = []
    if channel in ('msetuSrm', 'all'):
        sources.extend(_load_excel_data())
    if channel in ('poPortal', 'all'):
        from app.repositories.po_portal_invoice import get_all_cached_invoices
        sources.extend(get_all_cached_invoices())
    if not exclude_deleted and vendor_codes is None and not any(value is not None and value != '' for value in (
        status, vendor_code, invoice_number, po_number, po_item, product_type, plant, dv_type,
        upload_date_from, upload_date_to, filter_fiscal_year, date_from, date_to,
        approver, pending_with, approved_by, miro_status, payment_status, source_id,
        hod1_status, hod2_status, accounts_status, approval_stage, invoice_uploaded, po_type,
    )):
        return sources
    for row in sources:
        is_po = row.get('__channel') == 'poPortal'
        po_filters = [(hod1_status, 'HOD1_Status'), (hod2_status, 'HOD2_Status'),
                      (accounts_status, 'Accounts_Status'), (approval_stage, 'Status_OF_HOD'), (po_type, 'PO_Type')]
        if any(value is not None and (not is_po or (text_value(row.get(key)) or '').upper() != ('' if value == '__blank__' else value.strip().upper()))
               for value, key in po_filters):
            continue
        if invoice_uploaded is not None and ((text_value(row.get('Invoice_Uploaded')) or '').upper() == 'YES') != invoice_uploaded:
            continue
        if source_id is not None and row.get('__source_id') != source_id:
            continue
        vendor = (text_value(row.get("SUPPLIER")) or "").lower()
        if allowed is not None and vendor not in allowed:
            continue
        if vendor_code and vendor != vendor_code.strip().lower():
            continue
        if exclude_deleted and invoice_status(row) in ('Entry Deleted', 'Deleted'):
            continue
        approver_columns = ('Created_by_HOD1', 'Created_by_HOD2', 'Created_by_ACCOUNTS') if is_po else tuple(f'APPROVER{i}' for i in range(1, 8))
        if approver and not any(approver.strip().lower() in (text_value(row.get(column)) or '').lower() for column in approver_columns):
            continue
        if any(query and query.strip().lower() not in (text_value(row.get(column)) or '').lower()
               for query, column in [(pending_with, "PENDING_WITH"), (approved_by, "APPROVED_BY")]):
            continue
        miro = (text_value(row.get("MIRO_BOOKED")) or '').upper()
        if miro_status and miro != {"booked": "MIRO BOOKED", "pending": "PENDING FOR MIRO"}[miro_status]:
            continue
        if payment_status and (invoice_status(row) == "Paid") != (payment_status == "paid"):
            continue
        exact = ((invoice_number, "VENDOR_INVOICE_NO"), (po_number, "PURCHASING_DOCUMENT"),
                 (po_item, "ITEM"), (product_type, "PRODUCT_TYPE"), (plant, "PLANT"),
                 (filter_fiscal_year, "FILTER_FISCAL_YEAR"), (dv_type, "DV_TYPE"))
        if any(value is not None and str(value).strip() != "" and
               (text_value(row.get(column)) or "").lower() != str(value).strip().lower()
               for value, column in exact):
            continue
        if status:
            target = {"pending approval": "pending for approval", "deleted": "entry deleted"}.get(
                status.strip().lower(), status.strip().lower())
            if invoice_status(row).lower() != target:
                continue
        uploaded = date_value(row.get("INVOICE_UPLOAD_DATE"))
        if upload_date_from and (uploaded is None or uploaded < upload_date_from):
            continue
        if upload_date_to and (uploaded is None or uploaded > upload_date_to):
            continue
        if date_from or date_to:
            columns = DATE_FIELDS[date_type]
            if is_po:
                columns = {'miro': ('MIRO_POSTING_DATE',), 'posting': ('MIRO_POSTING_DATE',),
                           'approval': (), 'accounts': ()}.get(date_type, columns)
            selected_date = next((parsed for column in columns
                                  if (parsed := date_value(row.get(column))) is not None), None)
            if selected_date is None or (date_from and selected_date < date_from) or (date_to and selected_date > date_to):
                continue
        records.append(row)
    return records


def get_invoices(page=1, page_size=20, **filters):
    records = filter_invoices(**filters)
    start = (page - 1) * page_size
    return records[start:start + page_size], len(records)


def get_invoice_by_number(invoice_number: str, po_item: int | None = None, *, vendor_codes=None):
    records = filter_invoices(invoice_number=invoice_number, po_item=po_item, vendor_codes=vendor_codes)
    return records[0] if records else None


def get_all_cached_invoices():
    return _load_excel_data()
